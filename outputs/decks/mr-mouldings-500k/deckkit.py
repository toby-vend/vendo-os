"""Native PowerPoint primitives styled to the Vendo brand, for Google Slides import."""
from pptx import Presentation
from pptx.util import Inches as In, Pt, Emu
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

BG=C(0x05,0x14,0x12); CARD=C(0x09,0x22,0x1F); RAISED=C(0x0E,0x2C,0x28)
MINT=C(0x8E,0xFE,0xBB); WHITE=C(0xFF,0xFF,0xFF); BODY=C(0xE8,0xF2,0xEC)
SEC=C(0xB9,0xC7,0xC0); MUT=C(0x6B,0x7B,0x74); RULE=C(0x14,0x3A,0x35)
AMBER=C(0xF5,0xD6,0x74); RED=C(0xEF,0x44,0x44); MID=C(0x2C,0x5F,0x52)
SANS="Manrope"; SERIF="Instrument Serif"

W,H=13.333,7.5
ML=0.62; MR=0.62; CW=W-ML-MR

def new_deck():
    p=Presentation(); p.slide_width=In(W); p.slide_height=In(H); return p

def slide(prs):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    bg=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,In(W),In(H))
    bg.fill.solid(); bg.fill.fore_color.rgb=BG; bg.line.fill.background(); bg.shadow.inherit=False
    return s

def _tf(shape):
    tf=shape.text_frame; tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    return tf

def text(s,x,y,w,h,runs,size=14,color=BODY,bold=False,font=SANS,space=0,
         align=PP_ALIGN.LEFT,caps=False,line=None,anchor=MSO_ANCHOR.TOP):
    box=s.shapes.add_textbox(In(x),In(y),In(w),In(h)); tf=_tf(box); tf.vertical_anchor=anchor
    p=tf.paragraphs[0]; p.alignment=align
    if line: p.line_spacing=line
    if isinstance(runs,str): runs=[(runs,{})]
    for t,o in runs:
        r=p.add_run(); r.text=t.upper() if caps else t
        f=r.font; f.name=o.get("font",font); f.size=Pt(o.get("size",size))
        f.bold=o.get("bold",bold); f.italic=o.get("italic",False)
        f.color.rgb=o.get("color",color)
        sp=o.get("space",space)
        if sp: r.font._rPr.set('spc',str(int(sp*100)))
    return box

def para(s,x,y,w,h,lines,size=13,color=SEC,line=1.35,gap=8):
    """lines: list of run-lists (or strings) -> stacked paragraphs."""
    box=s.shapes.add_textbox(In(x),In(y),In(w),In(h)); tf=_tf(box)
    for i,ln in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.line_spacing=line; p.space_after=Pt(gap)
        if isinstance(ln,str): ln=[(ln,{})]
        for t,o in ln:
            r=p.add_run(); r.text=t; f=r.font
            f.name=o.get("font",SANS); f.size=Pt(o.get("size",size))
            f.bold=o.get("bold",False); f.italic=o.get("italic",False)
            f.color.rgb=o.get("color",color)
    return box

def bullets(s,x,y,w,h,items,size=13,color=SEC,gap=7):
    box=s.shapes.add_textbox(In(x),In(y),In(w),In(h)); tf=_tf(box)
    for i,it in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.line_spacing=1.32; p.space_after=Pt(gap)
        if isinstance(it,str): it=[(it,{})]
        it=[("•  ",{"color":MINT,"bold":True})]+list(it)
        for t,o in it:
            r=p.add_run(); r.text=t; f=r.font
            f.name=SANS; f.size=Pt(o.get("size",size)); f.bold=o.get("bold",False)
            f.color.rgb=o.get("color",color)
    return box

def card(s,x,y,w,h,fill=CARD,edge=RULE,radius=True):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                          In(x),In(y),In(w),In(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=fill
    sh.line.color.rgb=edge; sh.line.width=Pt(0.75); sh.shadow.inherit=False
    if radius:
        try: sh.adjustments[0]=0.045
        except Exception: pass
    sh.text_frame.text=""
    return sh

def rect(s,x,y,w,h,fill,edge=None):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,In(x),In(y),In(w),In(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if edge: sh.line.color.rgb=edge; sh.line.width=Pt(0.75)
    else: sh.line.fill.background()
    sh.shadow.inherit=False; return sh

def label(s,x,y,w,t,color=MUT,size=8.5):
    return text(s,x,y,w,0.2,t,size=size,color=color,bold=True,space=1.6,caps=True)

def stat(s,x,y,w,v,color=MINT,size=30):
    return text(s,x,y,w,0.55,v,size=size,color=color,bold=True,space=-0.6)

def pill(s,x,y,t,color=MINT,fill=None):
    w=0.11*len(t)+0.28
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,In(x),In(y),In(w),In(0.235))
    sh.fill.solid(); sh.fill.fore_color.rgb=fill or RAISED
    sh.line.fill.background(); sh.shadow.inherit=False
    try: sh.adjustments[0]=0.5
    except Exception: pass
    tf=sh.text_frame; tf.word_wrap=False
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
    r=p.add_run(); r.text=t.upper(); f=r.font
    f.name=SANS; f.size=Pt(7.5); f.bold=True; f.color.rgb=color
    f._rPr.set('spc','140')
    return sh

def table(s,x,y,w,headers,rows,widths,fs=10.5,rowh=0.29,hi=(),dim=(),rowcolors=None):
    n=len(rows)+1
    g=s.shapes.add_table(n,len(headers),In(x),In(y),In(w),In(0.3+rowh*len(rows))).table
    tot=sum(widths)
    for i,cw in enumerate(widths): g.columns[i].width=Emu(int(In(w)*cw/tot))
    g.rows[0].height=In(0.30)
    for i in range(1,n): g.rows[i].height=In(rowh)
    def style(cell,t,*,bold,color,size,align,fill):
        cell.fill.solid(); cell.fill.fore_color.rgb=fill
        cell.margin_left=In(0.05); cell.margin_right=In(0.05)
        cell.margin_top=In(0.02); cell.margin_bottom=In(0.02)
        cell.vertical_anchor=MSO_ANCHOR.MIDDLE
        tf=cell.text_frame; tf.word_wrap=True
        p=tf.paragraphs[0]; p.alignment=align
        r=p.add_run(); r.text=t; f=r.font
        f.name=SANS; f.size=Pt(size); f.bold=bold; f.color.rgb=color
    for c,htxt in enumerate(headers):
        style(g.cell(0,c),htxt.upper(),bold=True,color=MUT,size=7.8,
              align=PP_ALIGN.RIGHT if c and headers[c].startswith("~") is False and c>0 and _num(c,headers) else PP_ALIGN.LEFT,
              fill=BG)
    for r_i,row in enumerate(rows,start=1):
        base=BODY
        fill=BG
        if r_i-1 in hi: base=MINT; fill=C(0x0B,0x2A,0x24)
        if r_i-1 in dim: base=MUT
        if rowcolors and (r_i-1) in rowcolors: base=rowcolors[r_i-1]
        for c,val in enumerate(row):
            style(g.cell(r_i,c),str(val),bold=(r_i-1 in hi),color=base,size=fs,
                  align=PP_ALIGN.RIGHT if c>0 else PP_ALIGN.LEFT,fill=fill)
    return g

def _num(c,headers): return True

def footer(s,num,total=22):
    rect(s,ML,H-0.66,CW,0.011,RULE)
    text(s,ML,H-0.52,6,0.24,"MR MOULDINGS × VENDO DIGITAL",size=8,color=MUT,bold=True,space=1.8)
    text(s,W-MR-1.2,H-0.52,1.2,0.24,f"{num:02d}",size=8,color=MUT,bold=True,
         space=1.0,align=PP_ALIGN.RIGHT)

def header(s,eyebrow_t,title_runs,standfirst=None,num=None):
    label(s,ML,0.46,CW,eyebrow_t,color=MINT,size=9)
    text(s,ML,0.76,CW,0.7,title_runs,size=27,color=WHITE,bold=True,space=-0.7)
    y=1.52
    if standfirst:
        para(s,ML,y,CW-0.4,0.5,[standfirst],size=12,color=SEC,line=1.35,gap=0); y+=0.52
    if num: footer(s,num)
    return y+0.18

def flourish(t): return (t,{"font":SERIF,"italic":True,"color":MINT,"size":27})
