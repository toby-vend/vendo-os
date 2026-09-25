# Rebuilds each slide as native, editable PowerPoint shapes from the layout JSON
# captured by extract.js. Text stays text; only SVG charts/logos become images.
import glob, json, os, re
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
EMU = 12192000 / 1600          # EMU per CSS px
PT = 0.6                       # pt per CSS px (1600px == 960pt)
FONT = 0.955                   # Slides renders Manrope slightly wider than Chrome

def E(v): return Emu(int(round(v * EMU)))

def parse(c):
    m = re.match(r'rgba?\(([^)]+)\)', c or '')
    if not m: return None
    p = [float(x) for x in m.group(1).replace('/', ',').split(',') if x.strip()]
    a = p[3] if len(p) > 3 else 1.0
    return RGBColor(int(p[0]), int(p[1]), int(p[2])), a

def set_alpha(color_format_parent_xml, alpha):
    if alpha >= 0.999: return
    clr = color_format_parent_xml.find(qn('a:srgbClr'))
    if clr is None: return
    for old in clr.findall(qn('a:alpha')): clr.remove(old)
    el = clr.makeelement(qn('a:alpha'), {'val': str(int(alpha * 100000))})
    clr.append(el)

def fill_shape(shape, css):
    c = parse(css)
    if not c: shape.fill.background(); return
    rgb, a = c
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb
    set_alpha(shape.fill._xPr.find(qn('a:solidFill')), a)

def line_shape(shape, css, width_px):
    c = parse(css)
    if not c or width_px <= 0: shape.line.fill.background(); return
    rgb, a = c
    shape.line.color.rgb = rgb
    shape.line.width = Pt(max(width_px * PT, 0.5))
    ln = shape.line._get_or_add_ln()
    set_alpha(ln.find(qn('a:solidFill')), a)

def no_shadow(shape):
    sp = shape._element.spPr
    if sp.find(qn('a:effectLst')) is None:
        sp.append(sp.makeelement(qn('a:effectLst'), {}))

ALIGN = {'right': PP_ALIGN.RIGHT, 'end': PP_ALIGN.RIGHT, 'center': PP_ALIGN.CENTER}

def style_run(run, r):
    f = run.font
    f.name = r['font'] if r['font'] in ('Manrope', 'Instrument Serif') else 'Manrope'
    f.size = Pt(round(r['size'] * PT * FONT * 2) / 2)
    f.bold = r['weight'] >= 600
    f.italic = bool(r['italic'])
    c = parse(r['color'])
    rpr = run._r.get_or_add_rPr()
    if c:
        f.color.rgb = c[0]
        set_alpha(rpr.find(qn('a:solidFill')), c[1])
    if r.get('ls'): rpr.set('spc', str(int(round(r['ls'] * PT * 100))))
    for tag in ('a:ea', 'a:cs'):
        if rpr.find(qn(tag)) is None:
            rpr.append(rpr.makeelement(qn(tag), {'typeface': f.name}))

def fill_para(p, runs):
    for r in runs:
        if r.get('br'):
            p.add_line_break(); continue
        run = p.add_run(); run.text = r['t']; style_run(run, r)

def bulletise(p, indent_px, color):
    pPr = p._p.get_or_add_pPr()
    emu = int(indent_px * EMU)
    pPr.set('marL', str(emu)); pPr.set('indent', str(-emu))
    for tag in ('a:buClr', 'a:buFont', 'a:buChar', 'a:buNone'):
        for old in pPr.findall(qn(tag)): pPr.remove(old)
    c = parse(color)
    if c:
        buClr = pPr.makeelement(qn('a:buClr'), {})
        srgb = buClr.makeelement(qn('a:srgbClr'), {'val': str(c[0])})
        buClr.append(srgb); pPr.append(buClr)
    pPr.append(pPr.makeelement(qn('a:buChar'), {'char': '\u2022'}))

def add_text(slide, it):
    paras = it.get('paras') or [it['runs']]
    first = next(r for rs in paras for r in rs if not r.get('br'))
    x, y, w, h = it['x'], it['y'], it['w'], it['h']
    slack = 4
    al = ALIGN.get(it['align'], PP_ALIGN.LEFT)
    if al == PP_ALIGN.RIGHT: x -= slack
    elif al == PP_ALIGN.CENTER: x -= slack / 2
    w += slack
    tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(max(h, first['size'])))
    tf = tb.text_frame
    tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, rs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = al
        p.line_spacing = Pt(it['lh'] * PT)
        if it.get('paras') and it.get('gap'): p.space_after = Pt(it['gap'] * PT)
        if it.get('paras') and it.get('bullet'): bulletise(p, it.get('indent', 20), rs[0].get('color') if rs and not rs[0].get('br') else None)
        fill_para(p, rs)

def build(out_path, layout_dir=None, svg_dir=None):
    prs = Presentation(); prs.slide_width = Emu(12192000); prs.slide_height = Emu(6858000)
    layout_dir = layout_dir or os.path.join(HERE, 'layout'); svg_dir = svg_dir or os.path.join(HERE, 'svg')
    for f in sorted(glob.glob(os.path.join(layout_dir, '*.json'))):
        d = json.load(open(f)); base = os.path.basename(f)[:-5]
        s = prs.slides.add_slide(prs.slide_layouts[6])
        bg = parse(d['bg'])
        s.background.fill.solid(); s.background.fill.fore_color.rgb = bg[0] if bg else RGBColor(5, 20, 18)
        svg_i = 0
        for it in d['items']:
            k = it['k']
            if k == 'rect':
                shp_type = MSO_SHAPE.ROUNDED_RECTANGLE if it['radius'] > 0 else MSO_SHAPE.RECTANGLE
                sh = s.shapes.add_shape(shp_type, E(it['x']), E(it['y']), E(it['w']), E(it['h']))
                if it['radius'] > 0:
                    sh.adjustments[0] = min(0.5, it['radius'] / max(1, min(it['w'], it['h'])))
                fill_shape(sh, it['fill']); line_shape(sh, it['line'], it['lw']); no_shadow(sh)
                sh.text_frame.text = ''
            elif k == 'line':
                x1, y1 = it['x'], it['y']
                x2, y2 = (x1, y1 + it['h']) if it.get('vertical') else (x1 + it['w'], y1)
                ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
                c = parse(it['color'])
                ln.line.color.rgb = c[0]; ln.line.width = Pt(max(it['lw'] * PT, 0.5))
                set_alpha(ln.line._get_or_add_ln().find(qn('a:solidFill')), c[1])
            elif k == 'svg':
                png = os.path.join(svg_dir, f'{base}-svg{svg_i}.png'); svg_i += 1
                if os.path.exists(png):
                    s.shapes.add_picture(png, E(it['x']), E(it['y']), E(it['w']), E(it['h']))
            elif k == 'text':
                add_text(s, it)
    prs.save(out_path)
    return len(prs.slides)

if __name__ == '__main__':
    import sys
    if len(sys.argv) > 1:
        print(build(sys.argv[1], sys.argv[2], sys.argv[3]), 'slides ->', sys.argv[1])
    else:
        out = os.path.join(HERE, '..', 'export', 'MR Mouldings - Roadmap to 500k (editable).pptx')
        print(build(out), 'slides ->', os.path.normpath(out))
